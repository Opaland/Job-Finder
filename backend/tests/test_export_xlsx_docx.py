"""Contenu réel des exports xlsx/docx (issue #8) : jusqu'ici seule leur présence
en mode démo était vérifiée (test_demo_couverture.py), jamais ce qu'ils contiennent.
"""
import io
from types import SimpleNamespace

from openpyxl import load_workbook

from app.models import Offer
from app.services.exports_offres import classeur_xlsx, lettre_docx, nom_fichier_lettre


def _offre(db, **champs):
    base = {"fingerprint": f"fp-{champs.get('source_id', '1')}", "source": "apec",
            "source_id": "1", "title": "Test Manager", "company": "ACME",
            "final_score": 88.0, "status": "nouvelle"}
    base.update(champs)
    db.add(Offer(**base))
    db.commit()


def test_export_xlsx_contenu_et_tri_par_score(client, db):
    _offre(db, source_id="1", title="Moins bonne offre", final_score=40.0)
    _offre(db, source_id="2", title="Meilleure offre", final_score=90.0)

    resp = client.get("/api/offers/export.xlsx")
    assert resp.status_code == 200
    wb = load_workbook(io.BytesIO(resp.content))
    ws = wb.active

    entetes = [c.value for c in ws[1]]
    assert entetes[:3] == ["Score", "Titre", "Entreprise"]
    # Triée par score décroissant : la meilleure offre d'abord.
    assert ws.cell(row=2, column=2).value == "Meilleure offre"
    assert ws.cell(row=2, column=1).value == 90
    assert ws.cell(row=3, column=2).value == "Moins bonne offre"


def test_export_xlsx_neutralise_un_titre_qui_ressemble_a_une_formule(client, db):
    """Seule protection anti-injection du classeur Excel : `cell.data_type = "s"`
    force une lecture en texte plutôt qu'en formule à l'ouverture. Sans elle,
    un titre scrapé commençant par « = » s'exécuterait dans Excel."""
    _offre(db, source_id="inj", title="=1+1")
    resp = client.get("/api/offers/export.xlsx")
    wb = load_workbook(io.BytesIO(resp.content))
    ws = wb.active
    cellule = ws.cell(row=2, column=2)
    assert cellule.value == "=1+1"       # valeur intacte...
    assert cellule.data_type == "s"      # ...mais explicitement du texte, jamais une formule


def test_export_xlsx_emporte_lettre_prochaine_action_et_entretiens(client, db):
    """Issue #16 : l'export perdait le fil des relances et rendez-vous prévus —
    seuls score/titre/etc. étaient exportés, pas le suivi de candidature."""
    from datetime import datetime

    _offre(
        db, source_id="suivi", cover_letter="Madame, Monsieur, ma candidature...",
        next_action_date=datetime(2026, 11, 3), next_action_note="Relancer par mail",
        interviews=[{"date": "2026-10-20T10:00:00", "format": "visio",
                     "interlocuteur": "Jean Dupont", "ressenti": "bon"}],
    )
    resp = client.get("/api/offers/export.xlsx")
    wb = load_workbook(io.BytesIO(resp.content))
    ws = wb.active

    entetes = [c.value for c in ws[1]]
    assert "Prochaine action" in entetes and "Entretiens" in entetes and "Lettre" in entetes
    ligne = {entetes[i]: ws.cell(row=2, column=i + 1).value for i in range(len(entetes))}
    assert ligne["Lettre"] == "Madame, Monsieur, ma candidature..."
    assert ligne["Prochaine action"] == "03/11/2026 — Relancer par mail"
    assert "2026-10-20" in ligne["Entretiens"] and "Jean Dupont" in ligne["Entretiens"] and "bon" in ligne["Entretiens"]


def test_lettre_docx_404_si_offre_introuvable(client, db):
    assert client.get("/api/offers/999/letter.docx").status_code == 404


def test_lettre_docx_400_si_aucune_lettre(client, db):
    _offre(db, cover_letter="")
    offre = db.query(Offer).one()
    assert client.get(f"/api/offers/{offre.id}/letter.docx").status_code == 400


def test_lettre_docx_contenu_reel(client, db):
    from docx import Document

    _offre(db, title="Test Manager Santé", company="ACME Santé",
           cover_letter="Madame, Monsieur,\n\nPremier paragraphe de la lettre.\n\nCordialement.")
    offre = db.query(Offer).one()

    resp = client.get(f"/api/offers/{offre.id}/letter.docx")
    assert resp.status_code == 200
    assert "ACME_Sant" in resp.headers["content-disposition"]  # nom de fichier assaini

    doc = Document(io.BytesIO(resp.content))
    texte = "\n".join(p.text for p in doc.paragraphs)
    assert "Test Manager Santé" in texte and "ACME Santé" in texte
    assert "Premier paragraphe de la lettre." in texte
    assert "Cordialement." in texte


# --- Sprint 59 (issue #8) : fonctions pures testées sans passer par la pile HTTP ---

def _offre_legere(**champs):
    base = dict(
        final_score=70.0, title="Testeur QA", company="ACME", location="Lyon",
        contract_type="CDI", salary_text="", remote=False, status="nouvelle",
        favorite=False, source="apec", published_at=None, collected_at=None,
        ai_reason="", notes="", url="https://exemple.fr", next_action_date=None,
        next_action_note=None, interviews=[], cover_letter="",
    )
    base.update(champs)
    return SimpleNamespace(**base)


def test_classeur_xlsx_fonction_pure_sans_offre():
    """Un classeur vide ne doit pas planter (max(len(offers) + 1, 2))."""
    contenu = classeur_xlsx([])
    wb = load_workbook(io.BytesIO(contenu))
    assert [c.value for c in wb.active[1]][:2] == ["Score", "Titre"]


def test_lettre_docx_fonction_pure():
    from docx import Document

    offre = _offre_legere(title="QA Lead", company="Beta", cover_letter="Madame,\n\nTexte.\n\nCordialement.")
    profil = SimpleNamespace(full_name="Jean Dupont", email="jean@exemple.fr")

    contenu = lettre_docx(offre, profil)
    doc = Document(io.BytesIO(contenu))
    texte = "\n".join(p.text for p in doc.paragraphs)
    assert "Jean Dupont" in texte and "QA Lead" in texte and "Beta" in texte and "Texte." in texte


def test_nom_fichier_lettre_assainit_l_entreprise():
    offre = _offre_legere(company="ACME Santé !", id=42)
    assert nom_fichier_lettre(offre) == "lettre_ACME_Sant_42.docx"
