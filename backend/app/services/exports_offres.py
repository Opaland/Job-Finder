"""Construction des fichiers d'export (classeur Excel, lettre Word).

Sprint 59 (issue #8) : logique extraite de `routers/offers.py`, qui ne garde
plus que la requête DB et la construction de la `Response` HTTP — même patron
que `services/justificatif.py` (PDF) et `services/exports_offres.py` permet
désormais de tester le CONTENU généré sans passer par toute la pile HTTP.
"""
from __future__ import annotations

import io
import re

from ..models import STATUS_LABELS

# En-têtes du classeur — (libellé, largeur de colonne).
COLONNES_XLSX = [
    ("Score", 8), ("Titre", 42), ("Entreprise", 24), ("Lieu", 20), ("Contrat", 14),
    ("Salaire", 16), ("Télétravail", 11), ("Statut", 12), ("Favori", 8),
    ("Source", 16), ("Publiée le", 12), ("Collectée le", 12), ("Avis IA", 40),
    ("Notes", 40), ("Prochaine action", 30), ("Entretiens", 40), ("Lettre", 50), ("Lien", 45),
]


def classeur_xlsx(offers) -> bytes:
    """Classeur Excel listant `offers` (lignes triées par score), prêt à servir.

    `offers` : résultat de requête SQLAlchemy portant les colonnes utilisées
    ci-dessous (voir routers/offers.py:export_xlsx) — pas la description,
    trop lourde pour cet export.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Offres"
    header_fill = PatternFill("solid", fgColor="10192B")
    for col, (label, width) in enumerate(COLONNES_XLSX, start=1):
        cell = ws.cell(row=1, column=col, value=label)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = "A2"

    for row, o in enumerate(offers, start=2):
        prochaine_action = ""
        if o.next_action_date:
            prochaine_action = o.next_action_date.strftime("%d/%m/%Y")
            if o.next_action_note:
                prochaine_action += f" — {o.next_action_note}"
        entretiens = "\n".join(
            f"{(i.get('date') or '')[:10]} ({i.get('format') or '?'}"
            f"{', ' + i['interlocuteur'] if i.get('interlocuteur') else ''})"
            + (f" — {i['ressenti']}" if i.get('ressenti') else "")
            for i in (o.interviews or [])
        )
        values = [
            round(o.final_score), o.title, o.company, o.location, o.contract_type,
            o.salary_text, "Oui" if o.remote else "", STATUS_LABELS.get(o.status, o.status),
            "★" if o.favorite else "",
            o.source, o.published_at.strftime("%d/%m/%Y") if o.published_at else "",
            o.collected_at.strftime("%d/%m/%Y"), o.ai_reason, o.notes,
            prochaine_action, entretiens, o.cover_letter, o.url,
        ]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=row, column=col, value=value)
            if isinstance(value, str) and value.startswith("="):
                # `data_type` suffit ici, et ne modifie pas la valeur : dans un
                # .xlsx, l'apostrophe de `cellule_sure` (réservée au CSV, où
                # Excel la masque) s'afficherait telle quelle.
                cell.data_type = "s"  # texte collecté, jamais une formule Excel
            cell.alignment = Alignment(vertical="top", wrap_text=col in (2, 13, 14, 16, 17))
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLONNES_XLSX))}{max(len(offers) + 1, 2)}"

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def lettre_docx(offer, profile) -> bytes:
    """Lettre de motivation de `offer` au format Word, en-tête avec le profil."""
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    header = doc.add_paragraph()
    run = header.add_run(profile.full_name or "")
    run.bold = True
    if profile.email:
        header.add_run(f"\n{profile.email}")

    title = f"Candidature — {offer.title}"
    if offer.company:
        title += f" · {offer.company}"
    doc.add_paragraph(title).runs[0].bold = True
    doc.add_paragraph("")

    for paragraph in offer.cover_letter.split("\n\n"):
        if paragraph.strip():
            doc.add_paragraph(paragraph.strip())

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def nom_fichier_lettre(offer) -> str:
    """Nom de fichier assaini pour le téléchargement de la lettre Word."""
    safe_company = re.sub(r"[^A-Za-z0-9_-]+", "_", offer.company or "offre").strip("_") or "offre"
    return f"lettre_{safe_company}_{offer.id}.docx"
